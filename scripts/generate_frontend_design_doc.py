from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSET_DIR = DOCS / "generated" / "frontend-design-assets"
OUTPUT = DOCS / "具身智能机械臂教学平台_Web前端设计说明书.docx"

NAVY = "173B4F"
BLUE = "2E718F"
TEAL = "2F806B"
GOLD = "A9792B"
LIGHT_BLUE = "EAF4F7"
LIGHT_TEAL = "EAF4F1"
LIGHT_GOLD = "F7F2E7"
LIGHT_GRAY = "F3F5F6"
MID_GRAY = "D7E0E4"
TEXT_GRAY = "526873"
WHITE = "FFFFFF"


def set_cell_fill(cell, color):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_text(cell, text, bold=False, color=NAVY, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color="C8D4D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)
        borders.append(tag)


def add_table(document, headers, rows, widths=None, font_size=8.5):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        set_cell_fill(header.cells[index], NAVY)
        set_cell_text(header.cells[index], value, bold=True, color=WHITE, size=9)
        if widths:
            header.cells[index].width = Cm(widths[index])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            set_cell_fill(cells[index], WHITE if row_index % 2 == 0 else LIGHT_GRAY)
            set_cell_text(cells[index], value, size=font_size)
            if widths:
                cells[index].width = Cm(widths[index])
    document.add_paragraph()
    return table


def add_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(document, text, bold_prefix=None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.35
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        first.bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)
    return paragraph


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_numbered(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_callout(document, title, body, color=LIGHT_BLUE):
    table = document.add_table(rows=1, cols=1)
    set_table_borders(table, color=color, size="0")
    cell = table.cell(0, 0)
    set_cell_fill(cell, color)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    paragraph = cell.add_paragraph(body)
    paragraph.paragraph_format.space_after = Pt(3)
    document.add_paragraph()


def add_caption(document, text):
    paragraph = document.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.style = document.styles["Caption"]
    return paragraph


def add_page_break(document):
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def draw_box(ax, x, y, width, height, title, subtitle="", color="#EAF4F7"):
    rectangle = plt.Rectangle((x, y), width, height, facecolor=color, edgecolor="#9BB1BA", linewidth=1.2)
    ax.add_patch(rectangle)
    ax.text(x + width / 2, y + height * 0.62, title, ha="center", va="center", fontsize=10, weight="bold", color="#173B4F")
    if subtitle:
        ax.text(x + width / 2, y + height * 0.30, subtitle, ha="center", va="center", fontsize=7.2, color="#526873", wrap=True)


def arrow(ax, start, end, color="#51798B", style="->"):
    ax.annotate("", xy=end, xytext=start, arrowprops=dict(arrowstyle=style, color=color, lw=1.4))


def generate_architecture_diagram():
    plt.rcParams["font.sans-serif"] = ["DejaVu Sans"]
    fig, ax = plt.subplots(figsize=(13, 7.2), dpi=180)
    ax.set_xlim(0, 13)
    ax.set_ylim(0, 7.2)
    ax.axis("off")
    ax.text(0.3, 6.85, "Target Multi-user Teaching Platform Architecture", fontsize=17, weight="bold", color="#173B4F")
    ax.text(0.3, 6.52, "Browser experience, domain APIs, persistence and asynchronous execution", fontsize=9, color="#526873")
    draw_box(ax, 0.3, 5.25, 2.2, 0.9, "Students / Faculty / Admin", "Browser / HTTPS", "#F7F2E7")
    draw_box(ax, 3.0, 5.25, 2.6, 0.9, "React Web Shell", "Router · RBAC Route Guard · Query", "#EAF4F7")
    draw_box(ax, 6.1, 5.25, 2.4, 0.9, "Teaching Workspace", "Monaco / Three.js / ECharts", "#EAF4F7")
    draw_box(ax, 9.0, 5.25, 3.5, 0.9, "Admin & Analytics Workspace", "Tenant / Class / Provider / Dashboard", "#EAF4F1")
    arrow(ax, (2.5, 5.7), (3.0, 5.7))
    arrow(ax, (5.6, 5.7), (6.1, 5.7))
    arrow(ax, (8.5, 5.7), (9.0, 5.7))

    draw_box(ax, 0.3, 3.75, 2.2, 0.9, "Reverse Proxy / BFF", "TLS / Cookie / CSRF / Rate Limit", "#F3F5F6")
    draw_box(ax, 3.0, 3.75, 2.4, 0.9, "Identity & Tenant", "User / Org / Class / RBAC", "#F7F2E7")
    draw_box(ax, 5.8, 3.75, 2.4, 0.9, "Teaching Domain API", "Course / Project / Run / Eval", "#F7F2E7")
    draw_box(ax, 8.6, 3.75, 1.9, 0.9, "Agent Service", "Session / Patch / Audit", "#F7F2E7")
    draw_box(ax, 10.9, 3.75, 1.6, 0.9, "Analytics", "Events / Metrics", "#EAF4F1")
    for x in (1.4, 4.2, 7.0, 9.55, 11.7):
        arrow(ax, (x, 5.25), (x, 4.65))

    draw_box(ax, 0.3, 2.2, 2.2, 0.9, "PostgreSQL", "Transactional Data / Alembic", "#EAF4F1")
    draw_box(ax, 3.0, 2.2, 2.2, 0.9, "Object Storage", "Logs / Reports / Artifacts", "#EAF4F1")
    draw_box(ax, 5.7, 2.2, 1.8, 0.9, "Redis", "Cache / Session / Queue", "#EAF4F1")
    draw_box(ax, 8.0, 2.2, 2.0, 0.9, "Worker / Simulator", "Async Run / Mock / Adapter", "#F3F5F6")
    draw_box(ax, 10.5, 2.2, 2.0, 0.9, "Secret Manager", "Reference / Rotation / Audit", "#F3F5F6")
    for x in (1.4, 4.1, 6.6, 9.0, 11.5):
        arrow(ax, (x, 3.75), (x, 3.1))

    ax.add_patch(plt.Rectangle((0.3, 0.65), 12.2, 0.8, facecolor="#EAF4F1", edgecolor="#77A696", linewidth=1.2))
    ax.text(0.55, 1.08, "Cross-cutting", fontsize=10, weight="bold", color="#2F806B")
    ax.text(2.0, 1.08, "Structured Logs   Metrics & Traces   Audit   Retention   Backup   Alerts", fontsize=9, color="#315C50")
    ax.text(0.3, 0.24, "REST/JSON: commands & queries | WebSocket: telemetry | WebRTC: future simulation video | Signed URL: artifacts", fontsize=8.5, color="#526873")
    path = ASSET_DIR / "target_architecture.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def generate_wireframes():
    plt.rcParams["font.sans-serif"] = ["DejaVu Sans"]
    fig, axes = plt.subplots(3, 1, figsize=(12.5, 14), dpi=180)
    titles = ["Login & Organization Selection", "Course / Project List", "Experiment Workspace"]
    for ax, title in zip(axes, titles):
        ax.set_xlim(0, 12)
        ax.set_ylim(0, 6.2)
        ax.axis("off")
        ax.add_patch(plt.Rectangle((0.2, 0.25), 11.6, 5.45, facecolor="#FCFDFD", edgecolor="#9BB1BA", linewidth=1.4))
        ax.text(0.45, 5.9, title, fontsize=14, weight="bold", color="#173B4F")

    ax = axes[0]
    ax.add_patch(plt.Rectangle((0.2, 4.9), 11.6, 0.8, facecolor="#173B4F", edgecolor="none"))
    ax.text(0.55, 5.3, "Embodied Arm Studio", color="white", fontsize=12, weight="bold", va="center")
    draw_box(ax, 1.0, 1.0, 4.5, 3.3, "Platform Capabilities", "Course / Project / Simulation / Evaluation\nCapabilities and safety boundary", "#EAF4F7")
    ax.add_patch(plt.Rectangle((6.4, 1.0), 4.4, 3.3, facecolor="white", edgecolor="#9BB1BA"))
    ax.text(6.8, 3.85, "Sign in", fontsize=13, weight="bold", color="#173B4F")
    for y, label in ((3.15, "Email / Student ID"), (2.45, "Password")):
        ax.add_patch(plt.Rectangle((6.8, y - 0.32), 3.6, 0.5, facecolor="#F3F5F6", edgecolor="#C8D4D9"))
        ax.text(7.0, y - 0.05, label, fontsize=8, color="#526873")
    ax.add_patch(plt.Rectangle((6.8, 1.55), 3.6, 0.55, facecolor="#2E718F", edgecolor="none"))
    ax.text(8.6, 1.83, "Sign in & Select Org", ha="center", va="center", fontsize=9, color="white", weight="bold")
    ax.text(8.6, 1.25, "OIDC / Institutional SSO", ha="center", fontsize=7.5, color="#526873")

    ax = axes[1]
    ax.add_patch(plt.Rectangle((0.2, 4.9), 11.6, 0.8, facecolor="#173B4F", edgecolor="none"))
    ax.text(0.55, 5.3, "Courses & Projects", color="white", fontsize=12, weight="bold", va="center")
    ax.text(10.4, 5.3, "Org / Profile", color="white", fontsize=8, va="center")
    ax.add_patch(plt.Rectangle((0.2, 0.25), 2.1, 4.65, facecolor="#F3F5F6", edgecolor="none"))
    for index, label in enumerate(("Home", "My Courses", "Projects", "Runs", "Evaluation", "Admin")):
        ax.text(0.55, 4.45 - index * 0.55, label, fontsize=8.5, color="#173B4F")
    ax.text(2.7, 4.45, "My Courses", fontsize=13, weight="bold", color="#173B4F")
    ax.add_patch(plt.Rectangle((8.7, 4.18), 2.4, 0.48, facecolor="white", edgecolor="#C8D4D9"))
    ax.text(8.9, 4.42, "Search courses / projects", fontsize=7.5, color="#526873")
    for row in range(2):
        for col in range(2):
            x, y = 2.7 + col * 4.25, 2.55 - row * 1.55
            ax.add_patch(plt.Rectangle((x, y), 3.8, 1.25, facecolor="white", edgecolor="#C8D4D9"))
            ax.text(x + 0.25, y + 0.88, "Robot Arm RL Fundamentals", fontsize=9, weight="bold", color="#173B4F")
            ax.text(x + 0.25, y + 0.54, "Progress 4/8 / Last run succeeded", fontsize=7.5, color="#526873")
            ax.add_patch(plt.Rectangle((x + 2.55, y + 0.18), 0.95, 0.38, facecolor="#2E718F", edgecolor="none"))
            ax.text(x + 3.03, y + 0.37, "Continue", fontsize=7.5, color="white", ha="center", va="center")

    ax = axes[2]
    ax.add_patch(plt.Rectangle((0.2, 4.9), 11.6, 0.8, facecolor="#173B4F", edgecolor="none"))
    ax.text(0.55, 5.3, "Experiment: Reach Target", color="white", fontsize=11, weight="bold", va="center")
    ax.text(9.0, 5.3, "Stop  |  Start Run", color="white", fontsize=8, va="center")
    ax.add_patch(plt.Rectangle((0.2, 0.25), 2.35, 4.65, facecolor="#F3F5F6", edgecolor="none"))
    ax.text(0.5, 4.45, "Course Steps", fontsize=9.5, weight="bold", color="#173B4F")
    for index, label in enumerate(("1 Environment", "2 Observation", "3 Action", "4 Reward", "5 Run", "6 Evaluation")):
        ax.text(0.55, 3.95 - index * 0.52, label, fontsize=8, color="#173B4F")
    ax.add_patch(plt.Rectangle((2.75, 2.2), 5.55, 2.5, facecolor="#EAF4F7", edgecolor="#9BB1BA"))
    ax.text(5.52, 3.55, "Three.js Mock / WebRTC Video", ha="center", fontsize=11, weight="bold", color="#173B4F")
    ax.text(5.52, 3.10, "Robot, target, trajectory and status overlay", ha="center", fontsize=8, color="#526873")
    ax.add_patch(plt.Rectangle((2.75, 0.48), 5.55, 1.42, facecolor="white", edgecolor="#C8D4D9"))
    ax.text(3.0, 1.58, "Metrics & Run Events", fontsize=9, weight="bold", color="#173B4F")
    ax.text(3.0, 1.15, "Reward ───╮     Success ─────", fontsize=8, color="#2F806B")
    ax.text(3.0, 0.78, "STEP 80 · STATUS COMPLETED", fontsize=7.5, color="#526873")
    ax.add_patch(plt.Rectangle((8.55, 0.48), 3.0, 4.22, facecolor="white", edgecolor="#C8D4D9"))
    ax.text(8.85, 4.35, "Configuration & Evidence", fontsize=9.5, weight="bold", color="#173B4F")
    for index, label in enumerate(("Environment", "Strategy", "Seed", "Max steps", "Run status", "Evaluation")):
        ax.text(8.85, 3.85 - index * 0.5, label, fontsize=7.7, color="#526873")
        ax.add_patch(plt.Rectangle((9.8, 3.68 - index * 0.5), 1.4, 0.3, facecolor="#F3F5F6", edgecolor="#D7E0E4"))

    path = ASSET_DIR / "core_page_wireframes.png"
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def generate_data_model():
    plt.rcParams["font.sans-serif"] = ["DejaVu Sans"]
    fig, ax = plt.subplots(figsize=(13, 7), dpi=180)
    ax.set_xlim(0, 13)
    ax.set_ylim(0, 7)
    ax.axis("off")
    ax.text(0.3, 6.62, "Core Data Model and Tenant Boundary", fontsize=16, weight="bold", color="#173B4F")
    boxes = {
        "Organization": (0.4, 5.0), "User": (3.0, 5.0), "Membership": (5.6, 5.0), "Role/Permission": (8.2, 5.0),
        "Course/Class": (0.4, 3.2), "Enrollment": (3.0, 3.2), "Project": (5.6, 3.2), "Run": (8.2, 3.2), "Evaluation": (10.6, 3.2),
        "Agent Session": (3.0, 1.4), "Audit Event": (5.6, 1.4), "Metric Event": (8.2, 1.4), "Provider Config": (10.6, 1.4),
    }
    subtitles = {
        "Organization": "tenant_id · name", "User": "identity · status", "Membership": "org · user · role", "Role/Permission": "scope · action",
        "Course/Class": "owner · term", "Enrollment": "class · user · progress", "Project": "org · owner · course", "Run": "project · config · status", "Evaluation": "run · rubric · result",
        "Agent Session": "project · provider", "Audit Event": "actor · action · target", "Metric Event": "event · version · subject", "Provider Config": "secret_ref · model",
    }
    for name, (x, y) in boxes.items():
        draw_box(ax, x, y, 2.0, 0.82, name, subtitles[name], "#EAF4F7" if y > 4 else "#EAF4F1")
    relations = [
        ("Organization", "Membership"), ("User", "Membership"), ("Membership", "Role/Permission"),
        ("Organization", "Course/Class"), ("Course/Class", "Enrollment"), ("User", "Enrollment"),
        ("Course/Class", "Project"), ("Project", "Run"), ("Run", "Evaluation"),
        ("Project", "Agent Session"), ("Project", "Audit Event"), ("Run", "Metric Event"), ("Organization", "Provider Config"),
    ]
    for source, target in relations:
        sx, sy = boxes[source]
        tx, ty = boxes[target]
        arrow(ax, (sx + 1.0, sy), (tx + 1.0, ty + 0.82), color="#7895A2")
    ax.text(0.4, 0.45, "Business data is organization-scoped; repositories enforce tenant filters; audit events are append-only.", fontsize=8.5, color="#526873")
    path = ASSET_DIR / "core_data_model.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.line_spacing = 1.35

    for level, size, color in ((1, 18, NAVY), (2, 14, BLUE), (3, 11.5, TEAL)):
        style = document.styles[f"Heading {level}"]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(6)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("具身智能机械臂教学平台 · Web 前端与多用户平台基础设计说明书  |  ")
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def add_cover(document):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(90)
    run = paragraph.add_run("具身智能机械臂教学平台")
    run.bold = True
    run.font.size = Pt(30)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Web 前端与多用户平台基础设计说明书")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第一部分：补齐多用户 Web 平台基础")
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string(TEAL)
    document.add_paragraph()
    line = document.add_table(rows=1, cols=1)
    set_cell_fill(line.cell(0, 0), BLUE)
    line.cell(0, 0).height = Cm(0.12)
    document.add_paragraph()
    metadata = [
        ("文档性质", "前端架构与平台基础详细设计"),
        ("设计对象", "Embodied Arm Studio 教学研究网站"),
        ("依据版本", "本机半成品审计 + 全新多用户目标设计"),
        ("版本", "V1.0"),
        ("日期", "2026 年 8 月 16 日"),
        ("范围边界", "Web 平台、仿真教学与数据运营；不包含真实机械臂控制"),
    ]
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for key, value in metadata:
        cells = table.add_row().cells
        set_cell_fill(cells[0], LIGHT_BLUE)
        set_cell_text(cells[0], key, bold=True, size=10)
        set_cell_text(cells[1], value, size=10)
    document.add_paragraph()
    paragraph = document.add_paragraph("状态标记：已完成/可复用 · 部分完成/Mock · 本期新增 · 后续扩展")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.runs[0].font.size = Pt(9)
    paragraph.runs[0].font.color.rgb = RGBColor.from_string(TEXT_GRAY)
    add_page_break(document)


def add_toc(document):
    add_heading(document, "目录", 1)
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr_text, fld_sep, fld_end])
    add_body(document, "提示：在 Microsoft Word/WPS 中右键目录并选择“更新域”，即可生成页码。")
    add_page_break(document)


def build_document():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    architecture = generate_architecture_diagram()
    wireframes = generate_wireframes()
    data_model = generate_data_model()

    document = Document()
    configure_document(document)
    properties = document.core_properties
    properties.title = "具身智能机械臂教学平台 Web 前端与多用户平台基础设计说明书"
    properties.subject = "前端架构、页面交互、API、通信、多用户数据与运营指标"
    properties.author = "Embodied Arm Studio 项目组"
    properties.keywords = "React, FastAPI, RBAC, PostgreSQL, WebSocket, WebRTC, Dashboard"
    add_cover(document)
    add_toc(document)

    add_heading(document, "1. 文档目的与设计范围", 1)
    add_body(document, "本文给出具身智能机械臂教学研究平台第一阶段的 Web 前端和多用户平台基础设计。设计既复用本机半成品中已经验证的教学闭环，也允许按生产级网站重新划分架构，不以当前单体实现为限制。本文可作为前端开发、后端接口设计、数据库建模、测试验收和项目答辩的共同依据。")
    add_heading(document, "1.1 本期建设目标", 2)
    add_bullets(document, [
        "以正式关系型数据库和版本迁移替换进程内状态及分散 JSON 索引。",
        "建立用户登录、组织、课程班级和基于角色的访问控制（RBAC）。",
        "持久化项目、运行、评估、Agent 会话、Patch lineage 和审计记录。",
        "仅在服务端管理 Secret 和模型 Provider，浏览器不接触 API Key 明文。",
        "将单体 React 页面和 FastAPI 主模块按领域拆分，形成稳定接口边界。",
        "建立统一埋点、指标字典和运营 Dashboard，支持教学质量分析。",
    ])
    add_heading(document, "1.2 范围边界", 2)
    add_table(document, ["范围", "本期处理", "说明"], [
        ("多用户网站基础", "是", "认证、租户、班级、权限、数据库、审计、运营指标"),
        ("教学实验闭环", "复用并改造", "课程、项目、环境、Mock 训练、运行证据、评估、Agent"),
        ("高保真仿真视频", "接口预留", "目标采用 WebRTC，当前仍为 Three.js/Mock"),
        ("真实 GPU 强化学习", "非本期核心", "保留 Worker/Adapter 契约，不把 Mock 描述为真实训练"),
        ("真实机械臂控制", "不包含", "不建设 Robot Gateway、动作下发、实机标定或 E-stop 集成"),
    ], widths=[4.0, 3.0, 10.0])
    add_callout(document, "核心原则", "任何页面都必须把 ready、mock、planned、blocked 状态明确展示；缺少证据时不得把教学模拟写成真实训练或真实机器人能力。", LIGHT_GOLD)

    add_heading(document, "2. 本机现状与目标差距", 1)
    add_body(document, "本机项目位于 /home/user/embodied-arm-studio。代码审计显示，当前前端已具备课程、项目工作区、环境设计、训练、运行证据、评估、Pipeline、导出和教学 Agent 页面；FastAPI 已提供约 67 个 REST/WebSocket 路由。其主要问题不是“没有页面”，而是缺少多用户身份、正式数据模型、稳定领域边界和运营闭环。")
    add_table(document, ["能力", "当前状态", "可复用内容", "本期设计动作"], [
        ("前端框架", "已完成", "React 18 + TypeScript + Vite", "保留，拆分路由与领域模块"),
        ("工作台组件", "已完成", "Monaco、Three.js、ECharts、Lucide", "保留，建立统一 Design System"),
        ("数据请求", "已完成", "TanStack Query + fetch 封装", "增加认证续期、错误码、分页和租户头/上下文"),
        ("训练事件", "已完成/Mock", "WebSocket /ws/runs/{run_id}", "保留协议并加入鉴权、重连、游标和事件版本"),
        ("课程/项目/评估", "已完成基础", "核心页面和 API 可运行", "绑定 user_id、organization_id、class_id"),
        ("数据库", "待完成", "文件/目录和部分 JSON 存储", "PostgreSQL + Alembic Migration"),
        ("用户/RBAC", "待完成", "无正式登录和资源隔离", "OIDC/本地身份 + 组织级 RBAC"),
        ("Agent 状态", "部分内存", "会话、参与、Patch 流程", "会话、消息、审批与审计全部持久化"),
        ("Provider Secret", "部分完成", "服务端环境变量读取", "Secret 引用、加密、轮换和操作审计"),
        ("运营 Dashboard", "待完成", "ECharts 可复用", "埋点管道、指标字典、分群和看板"),
    ], widths=[3.0, 3.0, 5.0, 6.0], font_size=8)

    add_heading(document, "3. 总体技术架构", 1)
    document.add_picture(str(architecture), width=Inches(6.9))
    add_caption(document, "图 1 多用户教学研究 Web 平台目标架构")
    add_heading(document, "3.1 架构风格", 2)
    add_body(document, "第一阶段采用“模块化单体 + 独立异步 Worker”的务实架构。FastAPI 仍可作为一个部署单元，但认证、组织课程、项目实验、运行评估、Agent、Provider 和 Analytics 各自拥有 router、service、repository、schema 与权限策略。该方式比立即拆成微服务更适合当前团队规模，同时为未来按负载拆分 Worker、Analytics 或 Agent 服务保留边界。")
    add_heading(document, "3.2 前端技术栈及选型理由", 2)
    add_table(document, ["层次", "技术", "用途", "选择理由"], [
        ("核心框架", "React 18 + TypeScript + Vite", "页面与交互开发", "沿用现有代码；类型约束清晰；构建和生态成熟"),
        ("路由", "React Router 6/7（新增）", "嵌套路由、守卫、懒加载", "当前手写 pathname 分支不适合登录、组织和权限路由"),
        ("服务端状态", "TanStack Query", "缓存、重试、失效和异步状态", "现有项目已使用，适合 REST 查询与 mutation"),
        ("本地工作台状态", "Zustand", "当前课程步骤、选中策略、运行交互", "现有项目已使用；避免把服务端数据复制进全局状态"),
        ("管理 UI", "Ant Design 5（增量引入）", "表格、表单、抽屉、筛选、权限管理", "组织班级和 Dashboard 是高密度管理界面，组件成熟且可访问性较完整"),
        ("图标", "Lucide React", "命令按钮和状态图标", "现有项目已使用，视觉统一"),
        ("编辑器", "Monaco Editor", "训练代码、JSON 配置、Diff", "现有成熟编辑器组件，支持语法、诊断和差异查看"),
        ("三维预览", "Three.js", "轻量场景和 Mock 机械臂状态", "当前已集成；不与未来视频流职责混淆"),
        ("图表", "Apache ECharts", "训练指标和运营 Dashboard", "现有项目已使用，支持漏斗、留存、时序和大数据量"),
        ("测试", "Vitest + Testing Library + Playwright", "单元、组件和 E2E", "保留现有 Vitest/Playwright，补充组件可访问性测试"),
    ], widths=[2.5, 4.0, 4.0, 7.0], font_size=7.7)
    add_heading(document, "3.3 后端与基础设施", 2)
    add_bullets(document, [
        "FastAPI：提供版本化 REST API、WebSocket 运行事件和 OpenAPI 契约。",
        "PostgreSQL：保存用户、组织、班级、项目、运行索引、评估、Agent 会话、审计和 Provider 元数据。",
        "Alembic：数据库 Migration；每个版本包含 upgrade/downgrade，生产部署先备份后迁移。",
        "Redis：短期会话、速率限制、幂等键、实时事件游标和任务队列；不作为业务事实唯一来源。",
        "对象存储：保存运行日志、评估报告、截图、模型 Manifest 和其他大制品；数据库仅存元数据、校验和与 URI。",
        "Worker/Simulator Adapter：执行耗时训练或仿真任务，与 Web API 通过队列和事件协议解耦。",
        "Secret Manager：保存模型 Provider 密钥；数据库只保留 secret_ref、掩码和轮换时间。",
    ])

    add_heading(document, "4. 信息架构与路由设计", 1)
    add_table(document, ["一级模块", "主要路由", "目标用户", "主要内容"], [
        ("身份入口", "/login、/select-organization", "全部", "登录、MFA/SSO、组织选择、协议告知"),
        ("首页", "/home", "全部", "我的课程、最近项目、待处理任务、平台状态"),
        ("课程与班级", "/courses、/classes/:id", "学生/教师", "课程列表、班级成员、进度、作业与完成率"),
        ("项目工作区", "/projects/:id/*", "学生/研究者", "文件、环境、代码、快照、Agent、运行关联"),
        ("实验操作", "/projects/:id/experiments/:id", "学生/研究者", "仿真、训练配置、实时指标、停止与重试"),
        ("运行与评估", "/runs、/evaluations", "学生/教师", "证据、对比、评估报告、导出门禁"),
        ("教学 Agent", "/agent/sessions/:id", "学生/教师", "对话、设计前置检查、Patch 审批与回滚"),
        ("运营 Dashboard", "/admin/analytics", "教师/管理员", "访问、注册、活跃、课程、训练、评估、留存"),
        ("平台管理", "/admin/*", "管理员", "组织、用户、角色、Provider、审计与数据保留"),
    ], widths=[3.0, 4.5, 3.0, 7.0], font_size=8)
    add_heading(document, "4.1 路由守卫", 2)
    add_numbered(document, [
        "未登录用户访问受保护路由时跳转 /login，并保存安全的 return_to 相对路径。",
        "登录后若属于多个组织，进入 /select-organization；单组织用户直接进入 /home。",
        "前端根据服务端返回的 permission set 隐藏无权操作，但服务端仍对每次 API 请求重新鉴权。",
        "资源不存在和无权限统一采用不泄露资源信息的错误策略；跨组织 ID 不得通过修改 URL 访问。",
        "blocked/planned 能力在路由元数据中显式标识，按钮禁用并给出证据边界。",
    ])

    add_heading(document, "5. 核心页面布局与交互", 1)
    document.add_picture(str(wireframes), width=Inches(6.9))
    add_caption(document, "图 2 三个核心页面布局草图")
    add_heading(document, "5.1 登录与组织选择页", 2)
    add_body(document, "页面使用左右双区布局：左侧说明平台身份、能力边界和隐私信息；右侧为登录表单。点击“登录”调用 POST /api/v1/auth/login，成功后调用 GET /api/v1/me 和 GET /api/v1/organizations。若启用学校统一身份认证，则按钮跳转 GET /api/v1/auth/oidc/start，回调后由服务端写入 HttpOnly、Secure、SameSite Cookie。")
    add_bullets(document, [
        "表单区域：邮箱/学工号、密码、记住设备、登录、统一身份认证、忘记密码。",
        "错误处理：凭据错误不区分账号是否存在；连续失败触发速率限制和审计事件。",
        "组织选择：卡片显示组织名称、角色和最近访问时间；点击后调用 POST /api/v1/session/organization。",
        "跳转：登录成功 → 组织选择（多组织）→ /home；被邀请用户先进入邀请确认页。",
    ])
    add_heading(document, "5.2 课程/项目列表页", 2)
    add_body(document, "页面由顶部组织切换与用户菜单、左侧稳定导航、筛选工具条和内容列表组成。学生默认看到“我的课程”和最近项目，教师可以切换到班级管理视图。点击课程卡片调用 GET /api/v1/courses/{course_id} 并进入课程详情；点击“继续”按服务端返回的 next_step 跳转到项目或实验操作页。")
    add_bullets(document, [
        "筛选：学期、班级、状态、标签、负责人；搜索输入使用防抖并写入 URL query。",
        "课程卡：课程名、教师、进度、检查点、最近运行状态、下一步动作。",
        "项目卡：关联课程、环境版本、最近运行、评估状态和协作者。",
        "教师操作：创建班级、导入成员、发布课程、查看完成率；所有写操作进入审计。",
    ])
    add_heading(document, "5.3 实验操作页", 2)
    add_body(document, "实验页采用稳定的三栏工作台：左栏为课程步骤与完成条件，中部为仿真画面及训练曲线，右栏为配置、运行状态和证据。开始训练按钮先进行客户端结构校验，再调用 POST /api/v1/runs；成功后建立 WSS /api/v1/runs/{run_id}/events。停止按钮调用 POST /api/v1/runs/{run_id}/stop。运行完成后开放“评估”入口并调用 POST /api/v1/evaluations。")
    add_table(document, ["区域", "内容", "关键交互", "异常状态"], [
        ("顶部栏", "项目、环境版本、能力标识", "开始、停止、保存、返回", "无权限/运行冲突时禁用"),
        ("课程栏", "步骤、理论、预测、成功条件", "切换步骤、完成检查点", "未满足证据则不可完成"),
        ("仿真区", "Three.js 或 WebRTC 视频", "播放、暂停、重连、画质", "断流不终止后台任务"),
        ("指标区", "Reward、Success、Loss、事件", "缩放、筛选、导出 CSV", "断线后按 last_event_id 补拉"),
        ("配置栏", "环境、策略、seed、步数", "编辑、验证、启动", "422 显示字段级错误"),
        ("证据栏", "制品、评估、限制", "打开证据、比较运行", "缺失证据显示 not_assessed"),
    ], widths=[2.6, 4.0, 5.0, 5.5], font_size=8)

    add_heading(document, "6. 前端模块拆分设计", 1)
    add_body(document, "当前 apps/web/src/App.tsx 同时承担导航、页面、数据请求与业务交互。目标按“应用壳 + 领域 feature + 共享基础设施”拆分，页面只编排组件，不直接散布鉴权和错误处理。")
    add_table(document, ["目录", "职责", "示例"], [
        ("app/", "应用启动、Provider、Router、权限边界", "AppRouter.tsx、QueryProvider.tsx"),
        ("features/auth/", "登录、会话、组织切换", "LoginPage、useSession"),
        ("features/courses/", "课程、班级、进度", "CourseListPage、ClassRoster"),
        ("features/projects/", "项目、工作区、快照", "ProjectShell、WorkspaceEditor"),
        ("features/experiments/", "环境、训练、运行事件", "ExperimentPage、RunConsole"),
        ("features/evaluations/", "证据评估和对比", "EvaluationPage、EvidenceDrawer"),
        ("features/agent/", "会话、Patch 审批、回滚", "AgentPanel、PatchReview"),
        ("features/admin/", "组织、角色、Provider、Dashboard", "RoleMatrix、AnalyticsDashboard"),
        ("shared/api/", "生成客户端、错误、分页、事件协议", "client.ts、errors.ts"),
        ("shared/ui/", "主题、布局、状态和可访问组件", "PageHeader、CapabilityBadge"),
    ], widths=[4.5, 6.0, 6.5], font_size=8)
    add_heading(document, "6.1 状态与请求规则", 2)
    add_bullets(document, [
        "服务端事实由 TanStack Query 管理；Zustand 只保存未提交编辑状态和短期工作台交互。",
        "Query Key 必须包含 organizationId 及资源 ID，切换组织时清空不再可见的缓存。",
        "所有 mutation 使用 idempotency key；创建运行后禁止因按钮重复点击生成多个任务。",
        "API 客户端统一处理 401 刷新会话、403 权限、409 并发版本冲突、422 字段错误和 429 限流。",
        "路由级代码分割：课程、实验、评估、管理 Dashboard 分包加载；Monaco/Three.js/ECharts 仅在需要页面加载。",
    ])

    add_heading(document, "7. 前后端接口设计与调用清单", 1)
    add_body(document, "接口统一使用 /api/v1 前缀和 JSON。列表采用 cursor 分页；写操作返回 resource_version 或 ETag；删除默认软删除并写审计。下表覆盖图片所要求的登录、课程列表、训练、状态查询等前端调用，以及本阶段多用户和运营管理接口。")
    api_rows = [
        ("POST", "/auth/login", "登录页/点击登录", "建立服务端会话"),
        ("POST", "/auth/logout", "用户菜单/退出", "撤销会话并清缓存"),
        ("GET", "/me", "应用启动/路由恢复", "当前用户、角色、permissions"),
        ("GET", "/organizations", "登录后/组织切换", "可访问组织列表"),
        ("POST", "/session/organization", "点击组织", "切换当前租户上下文"),
        ("GET", "/courses", "课程列表加载/筛选", "课程、进度、next_step"),
        ("GET", "/courses/{id}", "点击课程卡", "课程详情与步骤"),
        ("GET", "/classes", "教师班级页", "班级及统计摘要"),
        ("POST", "/classes", "教师点击创建", "创建课程班级"),
        ("GET", "/classes/{id}/members", "打开成员页", "成员、角色、进度"),
        ("POST", "/classes/{id}/members", "导入/邀请成员", "建立 enrollment"),
        ("PUT", "/course-progress/{step}", "完成检查点", "保存进度和证据引用"),
        ("GET", "/projects", "项目列表加载", "按组织/课程过滤项目"),
        ("POST", "/projects", "点击新建项目", "初始化项目与工作区"),
        ("GET", "/projects/{id}", "进入项目", "项目元数据与权限"),
        ("GET", "/projects/{id}/workspace", "打开代码工作区", "文件树与版本"),
        ("PUT", "/projects/{id}/files/{path}", "保存文件", "乐观并发保存"),
        ("POST", "/projects/{id}/snapshots", "点击创建快照", "生成可回滚版本"),
        ("POST", "/projects/{id}/snapshots/{sid}/restore", "点击回滚", "恢复并写审计"),
        ("GET", "/environments", "实验页加载", "环境与 capability"),
        ("GET", "/strategies", "实验页加载", "策略与兼容性"),
        ("POST", "/runs", "点击开始训练", "创建幂等异步运行"),
        ("GET", "/runs/{id}", "刷新/重连/详情页", "运行状态与摘要"),
        ("POST", "/runs/{id}/stop", "点击停止", "请求取消运行"),
        ("GET", "/runs/{id}/events?after=", "WS 重连补拉", "按游标补齐事件"),
        ("WS", "/runs/{id}/events", "运行创建成功后", "状态、指标、错误事件"),
        ("GET", "/runs/{id}/artifacts", "运行详情/完成后", "制品元数据和签名 URL"),
        ("POST", "/evaluations", "点击评估", "按 run_ids 创建报告"),
        ("GET", "/evaluations/{id}", "评估详情", "证据、得分和限制"),
        ("POST", "/evaluations/compare", "点击报告对比", "基线与当前差异"),
        ("POST", "/agent/sessions", "首次打开 Agent", "创建持久化会话"),
        ("POST", "/agent/sessions/{id}/messages", "发送消息", "生成指导回复"),
        ("POST", "/agent/patches", "请求修改", "创建待审批 Diff"),
        ("POST", "/agent/patches/{id}/approve", "点击批准", "快照后应用 Patch"),
        ("GET", "/admin/roles", "角色管理页", "角色与权限矩阵"),
        ("PUT", "/admin/members/{id}/roles", "管理员保存角色", "更新 RBAC 并审计"),
        ("GET", "/admin/providers", "Provider 管理页", "配置状态，不返回密钥"),
        ("POST", "/admin/providers", "管理员保存 Provider", "写入 secret_ref"),
        ("POST", "/events/batch", "埋点缓冲触发", "批量上报版本化事件"),
        ("GET", "/admin/analytics/overview", "Dashboard 加载", "访问、注册、活跃摘要"),
        ("GET", "/admin/analytics/funnels", "课程/训练漏斗", "分群漏斗数据"),
        ("GET", "/admin/audit-events", "审计页查询", "按 actor/action/target 筛选"),
    ]
    add_table(document, ["方法", "接口（省略 /api/v1）", "页面/调用时机", "用途"], api_rows, widths=[1.8, 6.0, 5.0, 5.0], font_size=7.1)
    add_heading(document, "7.1 统一响应与错误规范", 2)
    add_body(document, "成功响应包含 data、meta、request_id；分页 meta 包含 next_cursor。错误响应采用 {code, message, field_errors, request_id}。前端只根据稳定 code 分支，不解析自然语言 message。推荐错误码包括 AUTH_REQUIRED、PERMISSION_DENIED、TENANT_MISMATCH、VERSION_CONFLICT、VALIDATION_FAILED、RUN_ALREADY_ACTIVE、RATE_LIMITED。")

    add_heading(document, "8. 仿真画面与实时通信方案", 1)
    add_heading(document, "8.1 当前实现", 2)
    add_body(document, "当前本机页面使用 Three.js 显示轻量机械臂和场景状态；训练创建后通过 WebSocket /ws/runs/{run_id} 接收 metric/status/error 事件。该链路已可用于 Mock 教学，但不等同于远程仿真视频。")
    add_heading(document, "8.2 目标双通道方案", 2)
    add_table(document, ["通道", "协议", "承载内容", "设计要求"], [
        ("命令/查询", "HTTPS REST", "创建运行、停止、配置、报告", "幂等、RBAC、审计、版本控制"),
        ("运行遥测", "WSS WebSocket", "metric、status、error、heartbeat", "JWT/Cookie 鉴权、事件序号、断线补拉"),
        ("仿真视频", "WebRTC", "H.264/VP8 视频，可选音频关闭", "STUN/TURN、拥塞控制、画质状态、重协商"),
        ("轻量状态", "Three.js + WSS", "关节、目标、轨迹、碰撞近似", "用于 Mock 或视频不可用降级，不冒充真实画面"),
        ("大制品", "HTTPS 签名 URL", "日志、报告、模型、录屏", "短时授权、校验和、下载审计"),
    ], widths=[2.5, 3.5, 5.5, 6.5], font_size=8)
    add_body(document, "视频会话流程：前端 POST /api/v1/simulations/{session_id}/stream-sessions 获取 stream_id、ICE servers 和权限；随后通过 WSS /api/v1/streams/{stream_id}/signal 交换 SDP offer/answer 与 ICE candidate；媒体轨道通过 WebRTC 传输。视频地址不设计成可公开复用的固定 URL，而是短时、用户与运行绑定的会话。运行遥测始终走独立 WebSocket，视频断线不会让训练状态丢失。")
    add_heading(document, "8.3 WebSocket 事件信封", 2)
    add_body(document, '{"event_id":"evt_...","sequence":128,"type":"metric","run_id":"run_...","occurred_at":"ISO-8601","schema_version":"1.0","payload":{"step":80,"metrics":{"mean_reward":0.82}}}')
    add_bullets(document, [
        "客户端保存最后 sequence；重连时带 after_sequence，服务端补发缺失事件。",
        "每 15 秒 heartbeat；连续两个周期无响应显示“实时连接中断”，但仍轮询 GET /runs/{id}。",
        "终态 completed/stopped/failed 到达后关闭连接并失效 runs、artifacts、evaluations 查询缓存。",
        "事件中不包含 Secret、用户敏感信息或宿主文件系统路径。",
    ])

    add_heading(document, "9. 数据库、Migration 与持久化", 1)
    document.add_picture(str(data_model), width=Inches(6.9))
    add_caption(document, "图 3 核心数据模型与租户边界")
    add_heading(document, "9.1 核心实体", 2)
    add_table(document, ["实体组", "主要表", "关键约束"], [
        ("身份租户", "users、organizations、memberships、roles、permissions", "身份唯一；成员关系唯一；角色带组织作用域"),
        ("教学组织", "courses、classes、enrollments、course_progress", "课程版本化；进度绑定用户和班级"),
        ("项目实验", "projects、workspace_versions、runs、run_artifacts", "项目归属组织；运行配置不可变；制品带 sha256"),
        ("评估", "evaluation_reports、criteria_results、evidence_refs", "rubric_version 固化；缺失证据为 not_assessed"),
        ("Agent", "agent_sessions、agent_messages、patches、patch_approvals", "会话归属项目；审批人与请求人可追溯"),
        ("平台配置", "provider_configs、secret_references、feature_flags", "永不存储可回显明文密钥"),
        ("运营审计", "metric_events、audit_events、retention_jobs", "事件版本化；审计只追加；按保留策略分区"),
    ], widths=[3.0, 7.0, 7.5], font_size=8)
    add_heading(document, "9.2 Migration 策略", 2)
    add_numbered(document, [
        "所有 schema 变更由 Alembic revision 管理，禁止应用启动时隐式建表。",
        "采用 expand → migrate data → switch read/write → contract 的兼容迁移顺序。",
        "部署前执行数据库备份和 migration dry-run；升级失败自动停止新版本，不继续启动 API。",
        "大表迁移采用后台回填与可恢复 checkpoint，避免长事务锁表。",
        "每个 revision 在 CI 中验证从上一生产版本 upgrade，并验证 downgrade 或明确不可逆说明。",
    ])
    add_heading(document, "9.3 现有数据迁移", 2)
    add_body(document, "现有 var/projects、var/runs、evaluations、pipelines 等目录先作为只读源导入。导入任务为每项生成 legacy_source、checksum 和 imported_at；重复执行保持幂等。大制品迁入对象存储，数据库保存索引。Agent 当前进程内会话无法可靠恢复，应在切换前明确标记为不可迁移，而不是伪造历史。")

    add_heading(document, "10. 登录、组织、课程班级与 RBAC", 1)
    add_heading(document, "10.1 身份与会话", 2)
    add_body(document, "优先支持学校 OIDC/OAuth 2.0 Authorization Code + PKCE；同时可保留受控本地账号用于开发。浏览器使用 HttpOnly Secure Cookie 持有会话标识，前端不把长期 token 写入 localStorage。状态修改请求使用 SameSite 策略和 CSRF token；会话可由管理员撤销。")
    add_heading(document, "10.2 角色权限矩阵", 2)
    add_table(document, ["能力", "Student", "Instructor", "Researcher", "Administrator"], [
        ("查看已加入课程", "允许", "允许", "按授权", "允许"),
        ("创建个人项目/运行", "允许", "允许", "允许", "允许"),
        ("查看他人项目", "禁止", "本班课程范围", "协作项目", "组织范围"),
        ("创建课程班级", "禁止", "允许", "可选", "允许"),
        ("查看班级完成率", "仅本人", "本班", "按授权", "组织范围"),
        ("审批高等级 Agent Patch", "按课程策略", "允许", "按项目策略", "允许"),
        ("管理 Provider", "禁止", "禁止", "禁止", "允许"),
        ("查看审计日志", "本人相关", "本班教学事件", "项目相关", "组织范围"),
        ("管理角色和数据保留", "禁止", "禁止", "禁止", "允许"),
    ], widths=[5.0, 3.0, 3.2, 3.2, 3.5], font_size=7.7)
    add_callout(document, "权限实现要求", "前端的菜单隐藏和按钮禁用只改善体验，不是安全边界。每个 repository 查询必须带 organization_id 作用域，每个 service 命令必须检查 permission、资源归属和当前课程策略。", LIGHT_GOLD)

    add_heading(document, "11. Secret 与模型 Provider 管理", 1)
    add_bullets(document, [
        "Provider 页面只显示名称、Base URL、可用模型、健康状态、密钥掩码后四位、更新时间和 secret_ref。",
        "管理员录入密钥时通过 HTTPS 直接提交服务端；服务端写入 Secret Manager 后立即丢弃请求体中的明文。",
        "普通日志、审计 payload、前端错误、数据库和导出文件都不得出现密钥明文。",
        "测试连接使用单独 POST /admin/providers/{id}/test，结果只返回 healthy/degraded/unavailable 和去敏错误。",
        "轮换采用“写入新版本 → 测试 → 切换引用 → 观察 → 撤销旧版本”，全过程写审计。",
        "Provider 使用受组织策略和数据范围约束；课程可指定允许的 Provider、模型和 Assistance Level。",
    ])

    add_heading(document, "12. 数据埋点与统一指标口径", 1)
    add_heading(document, "12.1 事件设计", 2)
    add_body(document, "前端埋点采用版本化事件，不直接把页面文案当事件名。每个事件包含 event_id、event_name、schema_version、occurred_at、anonymous_id/user_id、organization_id、session_id、course_id、project_id、run_id、app_version 和 properties。敏感代码、Agent 对话正文、Secret 和模型输入默认不进入运营埋点。")
    add_table(document, ["事件", "触发时机", "关键属性", "用途"], [
        ("page_viewed", "路由稳定后", "route_name、referrer", "访问量与页面路径"),
        ("registration_completed", "账号首次激活", "source、invite_id", "注册转化"),
        ("session_started", "建立有效会话", "device_class", "活跃与留存"),
        ("course_started", "首次进入课程步骤", "course_id、class_id", "课程漏斗"),
        ("course_step_completed", "证据满足并提交", "step_id、evidence_type", "学习进度"),
        ("course_completed", "全部必修步骤完成", "duration_seconds", "课程完成率"),
        ("run_requested", "创建运行请求", "environment、strategy", "训练启动量"),
        ("run_terminal", "运行进入终态", "status、reason、duration", "训练成功率"),
        ("evaluation_completed", "报告生成", "assessed、passed、rubric", "评估通过率"),
        ("agent_message_sent", "用户发送 Agent 消息", "assistance_level", "Agent 使用，不含正文"),
    ], widths=[4.2, 4.5, 5.0, 4.0], font_size=8)
    add_heading(document, "12.2 指标定义", 2)
    add_table(document, ["指标", "推荐口径", "说明"], [
        ("访问量 PV", "统计期内 page_viewed 数量", "排除健康检查、机器人和管理员测试流量"),
        ("独立访客 UV", "统计期内 distinct anonymous_id/user_id", "登录前后通过受控 identity merge 去重"),
        ("注册数", "registration_completed 的 distinct user_id", "邀请创建但未激活不计"),
        ("DAU/WAU/MAU", "有 qualifying_activity 的 distinct user_id", "仅登录不算活跃；至少课程、项目、运行或评估行为"),
        ("课程完成率", "完成课程人数 / 开始课程人数", "同一 course_version、class cohort 和统计窗口"),
        ("训练成功率", "completed runs / terminal runs", "queued/running 不进分母；用户主动 stopped 单独展示"),
        ("任务成功率", "达到环境 success termination 的 runs / assessed runs", "与系统执行成功率分开"),
        ("评估通过率", "passed reports / assessed reports", "not_assessed 不算失败，单列证据完整率"),
        ("D7 留存", "首个 qualifying_activity 后第 7 日回访用户 / cohort 用户", "建议同时给 D1、D7、D30"),
    ], widths=[3.4, 8.0, 6.3], font_size=8)

    add_heading(document, "13. Dashboard 设计", 1)
    add_body(document, "Dashboard 面向教师和管理员，不采用单一访问量大屏，而是以“获取 → 激活 → 学习 → 实验 → 评估 → 留存”漏斗组织。所有图表可按组织、班级、课程版本、用户角色、环境、策略和时间范围筛选，并显示数据更新时间与指标版本。")
    add_table(document, ["区域", "组件", "指标/维度", "点击后行为"], [
        ("总览", "KPI + 趋势", "PV、UV、注册、DAU/WAU/MAU", "进入对应明细与分群"),
        ("课程", "漏斗 + 班级排行", "开始、检查点、完成率、耗时", "查看掉队步骤和学生列表"),
        ("实验", "时序 + 失败原因", "运行量、执行成功率、任务成功率", "进入过滤后的 Runs"),
        ("评估", "通过率 + 完整性", "通过、not_assessed、证据缺失", "进入报告与缺失证据"),
        ("留存", "Cohort 热力图", "D1/D7/D30、课程 cohort", "下钻到课程版本"),
        ("Agent", "用量 + 教学辅助", "会话、消息、Assistance Level、Patch", "查看去敏聚合，不展示对话正文"),
        ("可靠性", "异常趋势", "API 错误、WS 断线、队列时长", "跳转可观测性系统"),
    ], widths=[2.5, 4.0, 6.0, 5.0], font_size=8)
    add_callout(document, "防止指标误导", "Dashboard 必须同时展示分子、分母、时间窗、过滤条件和指标版本。Mock 训练、真实训练（未来）和系统执行成功必须分栏，不能混成一个“成功率”。", LIGHT_GOLD)

    add_heading(document, "14. FastAPI 模块拆分与契约", 1)
    add_body(document, "当前 apps/api/main.py 直接注册大量路由并创建内存状态。目标保留 create_app 作为装配入口，将决策逻辑下沉到领域 service，将数据库操作收敛到 repository。")
    add_table(document, ["模块", "职责", "禁止事项"], [
        ("api/routers", "HTTP/WS 参数、鉴权依赖、状态码", "不直接访问数据库或文件系统"),
        ("domain/services", "用例、事务、权限和状态机", "不依赖 FastAPI Request"),
        ("domain/models", "领域实体、枚举、规则", "不包含传输层细节"),
        ("repositories", "PostgreSQL/对象存储访问", "不绕过 tenant scope"),
        ("schemas", "Pydantic 请求响应与事件版本", "不把 ORM 对象直接返回"),
        ("workers", "训练、仿真、评估异步任务", "不接受任意 shell 命令"),
        ("security", "会话、RBAC、CSRF、Secret", "不在日志输出凭据"),
        ("analytics", "事件校验、聚合、指标版本", "不采集代码/对话正文"),
    ], widths=[3.5, 7.5, 6.5], font_size=8)
    add_heading(document, "14.1 事务与一致性", 2)
    add_bullets(document, [
        "创建运行在同一事务中写 run 和 outbox event；Worker 从 outbox/queue 消费，避免数据库已提交但任务未发布。",
        "运行状态机仅允许 queued → running → completed/failed/stopped；非法逆转返回 409。",
        "文件和场景保存使用 resource_version；版本冲突由前端展示 Diff，不静默覆盖。",
        "评估只引用已终态运行及不可变 evidence_ref；报告保存 rubric_version 和计算版本。",
        "审计事件与高风险操作同事务写入；普通用户无删除审计权限。",
    ])

    add_heading(document, "15. 安全、隐私与可访问性", 1)
    add_table(document, ["主题", "设计要求"], [
        ("认证", "OIDC/短期会话、MFA 可选、HttpOnly Secure Cookie、会话撤销"),
        ("授权", "默认拒绝、组织作用域、对象级权限、服务端强制校验"),
        ("Web 安全", "CSRF、CORS allowlist、CSP、XSS 输出编码、SSRF/路径穿越防护、限流"),
        ("隐私", "目的限定、最小化采集、保留期限、用户导出/删除流程、去敏分析"),
        ("Secret", "只在服务端、密文引用、轮换、日志脱敏、最小权限"),
        ("Agent", "无 generic shell、allowlist 文件、Patch 必须展示/审批/快照/回滚"),
        ("可访问性", "键盘可操作、可见焦点、语义标签、图表文本摘要、颜色非唯一信号"),
        ("真实设备边界", "不提供真实机械臂动作路径；导出 Manifest 不启用控制"),
    ], widths=[3.2, 14.5], font_size=8.5)

    add_heading(document, "16. 性能、可靠性与测试", 1)
    add_heading(document, "16.1 前端性能目标", 2)
    add_bullets(document, [
        "普通管理页面首屏 JS 目标 gzip 后小于 350 KB；Monaco、Three.js、ECharts 独立懒加载。",
        "LCP 目标小于 2.5 秒、INP 小于 200 ms、CLS 小于 0.1（实验室典型网络）。",
        "列表默认 cursor 分页；运行事件图表采用窗口化/降采样，避免长训练耗尽浏览器内存。",
        "WebSocket 断线指数退避，恢复后补拉；页面刷新不会取消后台运行。",
    ])
    add_heading(document, "16.2 测试矩阵", 2)
    add_table(document, ["层次", "重点用例", "通过标准"], [
        ("Unit", "指标公式、RBAC 策略、运行状态机、事件校验", "关键规则分支全覆盖"),
        ("Component", "登录表单、课程筛选、权限按钮、错误态", "交互和可访问性断言通过"),
        ("API Contract", "OpenAPI、错误码、分页、幂等、租户隔离", "前后端生成类型一致"),
        ("Migration", "上一版本升级、回滚、数据导入幂等", "无丢失、可恢复"),
        ("Integration", "PostgreSQL、对象存储、Redis、Worker", "事务和故障路径通过"),
        ("Playwright E2E", "登录→课程→项目→训练→评估", "Student/Instructor/Admin 三角色通过"),
        ("Security", "越权、CSRF、CORS、SSRF、路径穿越、Secret 泄漏", "无高危问题"),
        ("Reliability", "WS 断线、Worker 重启、数据库故障", "可重连、任务状态一致"),
    ], widths=[3.2, 9.0, 5.5], font_size=8)

    add_heading(document, "17. 实施顺序与里程碑", 1)
    add_table(document, ["阶段", "交付内容", "退出条件"], [
        ("M1 数据与身份底座", "PostgreSQL/Alembic、User/Org/Membership、登录会话", "迁移可回滚；多用户隔离测试通过"),
        ("M2 课程与项目归属", "班级、Enrollment、项目/工作区租户化", "教师与学生权限流程通过"),
        ("M3 实验数据持久化", "Run、Evaluation、Agent、Audit、制品索引", "服务重启不丢关键状态"),
        ("M4 前后端模块拆分", "Router、feature 目录、领域 API、统一错误", "现有教学闭环回归通过"),
        ("M5 Provider 与 Secret", "Provider 管理、Secret 引用、轮换", "浏览器/日志/数据库无明文 Key"),
        ("M6 Analytics", "事件 SDK、指标字典、聚合和 Dashboard", "七类核心指标可追溯到版本化事件"),
        ("M7 稳定性验收", "E2E、安全、性能、备份恢复", "本章验收标准全部满足"),
    ], widths=[3.0, 8.5, 6.2], font_size=8)

    add_heading(document, "18. 验收标准", 1)
    add_numbered(document, [
        "两个不同组织中的同名项目和运行互不可见，直接修改 URL/API ID 也无法越权。",
        "数据库从上一版本迁移到目标版本成功，失败时可恢复；关键业务状态在 API 重启后保持。",
        "Student、Instructor、Researcher、Administrator 的角色矩阵由自动化测试验证。",
        "登录页、课程列表页、实验操作页完成桌面端和窄屏适配，键盘与焦点操作可用。",
        "训练创建具备幂等性；WebSocket 断线可重连补拉，终态与数据库一致。",
        "Agent 会话、消息、Patch、审批、快照、回滚和参与记录全部可审计。",
        "Provider 密钥不出现在浏览器响应、数据库明文字段、应用日志、审计 payload 或导出文件。",
        "PV/UV、注册、DAU/WAU/MAU、课程完成率、训练执行成功率、任务成功率、评估通过率和 D1/D7/D30 留存均有口径文档。",
        "Dashboard 每个指标可追溯到事件版本、查询版本、分子、分母和过滤条件。",
        "现有课程→项目→环境→训练→证据→评估→导出教学流程回归通过，Mock 状态持续明确展示。",
    ])

    add_heading(document, "附录 A：需求覆盖对照", 1)
    add_table(document, ["用户要求", "对应章节", "设计结论"], [
        ("正式数据库与 Migration", "第 9 章", "PostgreSQL + Alembic，兼容迁移和回滚"),
        ("登录、组织、课程班级、RBAC", "第 5、10 章", "OIDC/会话、租户和角色矩阵"),
        ("项目、运行、评估、Agent、审计持久化", "第 9、14 章", "事务数据、对象存储和审计一致性"),
        ("Secret 和 Provider 管理", "第 11 章", "服务端 Secret 引用、轮换和去敏"),
        ("拆分单体前端和 FastAPI", "第 6、14 章", "feature 前端 + 模块化单体后端"),
        ("数据埋点与统一口径", "第 12 章", "版本化事件和指标定义"),
        ("运营 Dashboard", "第 13 章", "获取、激活、学习、实验、评估、留存"),
        ("图片中的技术栈与理由", "第 3.2 节", "逐项说明框架和组件选型"),
        ("三个核心页面布局草图", "第 5 章", "登录、课程/项目、实验操作"),
        ("前后端接口清单", "第 7 章", "页面、按钮、调用时机和用途"),
        ("仿真画面接收方式", "第 8 章", "Three.js 当前；WebRTC 视频 + WS 遥测目标"),
    ], widths=[6.0, 3.0, 8.0], font_size=8)

    add_heading(document, "附录 B：设计依据与事实声明", 1)
    add_bullets(document, [
        "本机现有前端：React 18、TypeScript、Vite、TanStack Query、Zustand、Monaco、Three.js、ECharts、Lucide。",
        "本机现有后端：FastAPI 单体入口，包含课程、项目工作区、环境、训练、运行、评估、Agent、Sandbox 与 WebSocket 路由。",
        "本机当前边界：无正式关系型数据库、正式登录/RBAC、运营埋点和生产 WebRTC；部分 Agent 状态仍在内存。",
        "本文中的 PostgreSQL、Alembic、React Router、Ant Design、Redis、Secret Manager、WebRTC 和 Dashboard 属于目标设计，不表示已经实现。",
        "参考附件《设计书C》在会话中未提供可读取正文，仅可确认其为 Word 文档；本文未引用或虚构其中的具体段落，而是按用户要求补充完整架构和责任边界。",
    ])

    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = build_document()
    print(output)