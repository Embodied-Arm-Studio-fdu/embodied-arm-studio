#!/usr/bin/env python3
"""Idempotently document the three production manipulation projects."""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


DOCUMENT = Path(__file__).resolve().parents[1] / "docs" / "design_v1.docx"


def insert_paragraph_before(paragraph, text, style):
    node = OxmlElement("w:p")
    paragraph._p.addprevious(node)
    created = Paragraph(node, paragraph._parent)
    created.style = style
    created.add_run(text)
    return created


def main():
    document = Document(DOCUMENT)
    catalog = next(
        table for table in document.tables
        if table.cell(0, 0).text.strip() == "environment_type"
    )
    existing = {row.cells[0].text.strip() for row in catalog.rows[1:]}
    rows = (
        ("s1_grasp", "training", "对齐、下降、闭合夹爪与抓取验证；可调 XY 随机化和动作步长；输出 checkpoint、抓取指标、实时视频与录像。"),
        ("s1_lift", "training", "从稳定抓取状态抬升到 60 mm 目标带；可调动作步长和 1/5 秒保持；输出 checkpoint、抬升/稳定指标、实时视频与录像。"),
        ("s1_lower_release", "training", "下降到合法释放高度、策略触发松爪并验证稳定放置；可调释放阈值；输出 checkpoint、放置指标、实时视频与录像。"),
    )
    for task_id, kind, description in rows:
        if task_id in existing:
            continue
        cells = catalog.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = task_id, kind, description

    replacements = {
        "当前正式平台已接通 S1 Push、S1 Reach 两个 Training 和五项 S1 控制 Demo。":
            "当前正式平台已接通 S1 Push、S1 Reach、S1 Grasp、S1 Grasp & Lift、S1 Lower & Release 五个 Training 和五项 S1 控制 Demo。",
        "网页通过 Backend 认证代理读取七任务 catalog；Basic 提供 S1 Push Short + 5 Demo，Push/Reach Standard 各自只提供匹配 Training。":
            "网页通过 Backend 认证代理读取十任务 catalog；原有 Basic、Push 与 Reach 入口保持不变，并新增 Grasp、Lift、Lower & Release 三个只允许匹配 Project 的 Standard Training。",
        "Simulation registry 的 7 项任务均已接入正式 Backend/Run 链，但仍受 Project、run_kind、catalog 参数和权限约束。":
            "Simulation registry 的 10 项任务均已接入正式 Backend/Run 链，但仍受 Project、run_kind、catalog 参数和权限约束。",
        "三个 Project 均沿用 canonical Run：名称按任务名加 Run #ID 展示，配置由 Project 版本与 task catalog 冻结；":
            "三个 Project 均沿用 canonical Run：课程页使用“抓取训练、抓取抬起、抓取放下”三个简洁名称，不在实验名称中拼接任务 ID 或 Project ID；配置由 Project 版本与 task catalog 冻结；",
    }
    for paragraph in document.paragraphs:
        for old, new in replacements.items():
            if old in paragraph.text:
                paragraph.text = paragraph.text.replace(old, new)

    if not any(p.text.strip() == "10.5 三项 Manipulation 训练合同" for p in document.paragraphs):
        anchor = next(p for p in document.paragraphs if p.text.strip().startswith("11. Workspace"))
        insert_paragraph_before(anchor, "10.5 三项 Manipulation 训练合同", "Heading 2")
        insert_paragraph_before(
            anchor,
            "s1_grasp 复用已训练的 GraspPPOEnvironment 与四维 XYZ+夹爪动作；主要可视化指标包括 candidate_rate、verified_success_rate、对齐/下降率、接触阻塞和安全失败。s1_lift 从稳定抓取 reset 开始，以单轴 Z 动作进入 55–65 mm 目标带，显示 mean_maximum_cube_rise、mean_hold_duration、相对平移/旋转漂移和稳定抓取证据。s1_lower_release 从已抓取且位于目标上方的状态开始，使用下降与 release trigger，显示 descended_rate、release_latched_rate、detached_rate、settled_1s_rate 和 mean_final_xy_error。",
            "Normal",
        )
        insert_paragraph_before(
            anchor,
            "三个 Project 均沿用 canonical Run：课程页使用“抓取训练、抓取抬起、抓取放下”三个简洁名称，不在实验名称中拼接任务 ID 或 Project ID；配置由 Project 版本与 task catalog 冻结；网页仅允许 catalog 白名单参数，提供一键启动、停止、实时视频开关、终态 MP4 回放、指标 CSV、JSON 证据包及逐个 Artifact 下载。视频关闭只影响观察与录像，不停止训练；真实机器人控制保持关闭。",
            "Normal",
        )
        insert_paragraph_before(
            anchor,
            "当前三个 Manipulation 任务以各自训练环境产生的验证/稳定性指标和 checkpoint 作为任务证据，run_mode 使用 smoke 评估边界，不调用只支持 Push/Reach 的通用 100-episode ProductionEvaluationLauncher，也不伪造 evaluation_summary。Run completed 仅表示训练与必需制品完成；策略是否成功必须检查对应 verified/hold/settled 指标。",
            "Normal",
        )

    temporary = DOCUMENT.with_name(".design_v1.manipulation.tmp.docx")
    document.save(temporary)
    temporary.replace(DOCUMENT)


if __name__ == "__main__":
    main()
